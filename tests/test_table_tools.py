import asyncio
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import TestCase
from unittest.mock import patch
import zipfile

from openpyxl import Workbook, load_workbook
from starlette.datastructures import UploadFile
from starlette.requests import Request

from app.routers.api import api_download_vulnerability_result, api_vulnerability_process
from app.services.table_tools import TableInput, process_vulnerability_files


def _write_workbook(path: Path, headers: list[str], rows: list[list[object]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "数据"
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    workbook.save(path)


def _workbook_upload(filename: str, headers: list[str], rows: list[list[object]]) -> UploadFile:
    buffer = BytesIO()
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    workbook.save(buffer)
    buffer.seek(0)
    return UploadFile(buffer, filename=filename)


class VulnerabilityProcessingTests(TestCase):
    def test_builds_final_hss_report_with_ip_marking_filtering_and_deduplication(self) -> None:
        with TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            first_hss = root / "hss-1.xlsx"
            second_hss = root / "hss-2.xlsx"
            elb = root / "elb.xlsx"
            output = root / "final.xlsx"
            hss_headers = ["漏洞ID", "风险等级", "服务器IP"]
            _write_workbook(
                first_hss,
                hss_headers,
                [
                    ["V-1", "高危", "10.0.0.1"],
                    ["V-1", "高危", "10.0.0.1"],
                    ["V-2", "低危", "10.0.0.2"],
                    ["V-3", "中危", "10.0.0.9"],
                ],
            )
            _write_workbook(second_hss, hss_headers, [["V-4", "严重", "10.0.0.2"]])
            _write_workbook(
                elb,
                ["后端服务器-私网IP地址", "负载均衡"],
                [
                    ["10.0.0.1", "ELB-A"],
                    ["10.0.0.1", "ELB-A"],
                    ["10.0.0.2", "ELB-B"],
                ],
            )

            stats = process_vulnerability_files(
                [TableInput(first_hss, first_hss.name), TableInput(second_hss, second_hss.name)],
                [TableInput(elb, elb.name)],
                output,
                remove_risk_levels=["低危"],
            )

            self.assertEqual(stats["hss_rows_read"], 5)
            self.assertEqual(stats["hss_duplicate_rows"], 1)
            self.assertEqual(stats["filtered_rows"], 1)
            self.assertEqual(stats["filtered_by_level"], {"低危": 1})
            self.assertEqual(stats["rows_written"], 3)
            self.assertEqual(stats["elb_rows_read"], 3)
            self.assertEqual(stats["elb_duplicate_rows"], 1)
            self.assertEqual(stats["elb_unique_ips"], 2)
            self.assertEqual(stats["public_rows"], 2)
            self.assertEqual(stats["non_public_rows"], 1)

            workbook = load_workbook(output, read_only=True, data_only=True)
            try:
                rows = list(workbook["合并数据"].iter_rows(values_only=True))
            finally:
                workbook.close()
            self.assertEqual(rows[0], ("漏洞ID", "风险等级", "服务器IP", "是否是对外服务主机"))
            self.assertEqual(rows[1], ("V-1", "高危", "10.0.0.1", "是"))
            self.assertEqual(rows[2], ("V-3", "中危", "10.0.0.9", "否"))
            self.assertEqual(rows[3], ("V-4", "严重", "10.0.0.2", "是"))

    def test_processes_hss_and_elb_workbooks_from_zip_files(self) -> None:
        with TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            hss_workbook = root / "hss.xlsx"
            elb_workbook = root / "elb.xlsx"
            hss_archive = root / "hss.zip"
            elb_archive = root / "elb.zip"
            output = root / "final.xlsx"
            _write_workbook(
                hss_workbook,
                ["漏洞ID", "风险等级", "服务器IP"],
                [["V-1", "高危", "10.0.0.1"]],
            )
            _write_workbook(
                elb_workbook,
                ["后端服务器-私网IP地址"],
                [["10.0.0.1"]],
            )
            with zipfile.ZipFile(hss_archive, "w", zipfile.ZIP_DEFLATED) as archive:
                archive.write(hss_workbook, "华南/HSS漏洞报告.xlsx")
            with zipfile.ZipFile(elb_archive, "w", zipfile.ZIP_DEFLATED) as archive:
                archive.write(elb_workbook, "广州/ELB后端.xlsx")

            stats = process_vulnerability_files(
                [TableInput(hss_archive, hss_archive.name)],
                [TableInput(elb_archive, elb_archive.name)],
                output,
            )

            self.assertEqual(stats["hss_workbooks"], 1)
            self.assertEqual(stats["elb_workbooks"], 1)
            self.assertEqual(stats["rows_written"], 1)
            self.assertEqual(stats["public_rows"], 1)

    def test_rejects_hss_report_without_required_server_ip_header(self) -> None:
        with TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            hss = root / "hss.xlsx"
            elb = root / "elb.xlsx"
            _write_workbook(hss, ["漏洞ID", "风险等级"], [["V-1", "高危"]])
            _write_workbook(elb, ["后端服务器-私网IP地址"], [["10.0.0.1"]])

            with self.assertRaisesRegex(ValueError, "服务器IP"):
                process_vulnerability_files(
                    [TableInput(hss, hss.name)],
                    [TableInput(elb, elb.name)],
                    root / "result.xlsx",
                )

    def test_rejects_elb_report_without_backend_ip_header(self) -> None:
        with TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            hss = root / "hss.xlsx"
            elb = root / "elb.xlsx"
            _write_workbook(hss, ["漏洞ID", "风险等级", "服务器IP"], [["V-1", "高危", "10.0.0.1"]])
            _write_workbook(elb, ["IP地址"], [["10.0.0.1"]])

            with self.assertRaisesRegex(ValueError, "后端服务器-私网IP地址"):
                process_vulnerability_files(
                    [TableInput(hss, hss.name)],
                    [TableInput(elb, elb.name)],
                    root / "result.xlsx",
                )


class VulnerabilityProcessingApiTests(TestCase):
    def test_upload_process_and_download_flow(self) -> None:
        with TemporaryDirectory() as temp_dir:
            export_dir = Path(temp_dir)
            hss_upload = _workbook_upload(
                "HSS漏洞报告.xlsx",
                ["漏洞ID", "风险等级", "服务器IP"],
                [["V-1", "高危", "10.0.0.1"], ["V-2", "低危", "10.0.0.2"]],
            )
            elb_upload = _workbook_upload(
                "ELB后端.xlsx",
                ["后端服务器-私网IP地址"],
                [["10.0.0.1"]],
            )
            request = Request({"type": "http", "method": "POST", "path": "/api/tools/vulnerability-process", "headers": []})

            with (
                patch("app.routers.api.EXPORT_DIR", export_dir),
                patch(
                    "app.routers.api.require_login",
                    return_value={"username": "tester", "display_name": "测试人员"},
                ),
            ):
                result = asyncio.run(
                    api_vulnerability_process(
                        request,
                        hss_files=[hss_upload],
                        elb_files=[elb_upload],
                        remove_risk_levels="低危",
                    )
                )
                response = asyncio.run(
                    api_download_vulnerability_result(request, result["job_id"])
                )

            self.assertEqual(result["stats"]["rows_written"], 1)
            self.assertEqual(result["stats"]["filtered_rows"], 1)
            self.assertEqual(result["stats"]["public_rows"], 1)
            self.assertEqual(
                Path(response.path),
                export_dir / "vulnerability-tools" / result["job_id"] / "result.xlsx",
            )
            self.assertRegex(result["filename"], r"^比亚迪项目主机安全体检报告_\d{4}年\d{1,2}月\.xlsx$")

            workbook = load_workbook(response.path, read_only=True, data_only=True)
            try:
                rows = list(workbook["合并数据"].iter_rows(values_only=True))
            finally:
                workbook.close()
            self.assertEqual(rows[1], ("V-1", "高危", "10.0.0.1", "是"))
