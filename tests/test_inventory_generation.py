from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import TestCase
from unittest.mock import patch

from openpyxl import load_workbook

from app.services import inventory


class InventoryGenerationTests(TestCase):
    def test_build_match_key_extracts_ip_from_annotated_or_multi_ip_text(self) -> None:
        self.assertEqual(
            inventory.build_match_keys("", "10.167.89.35（私有）", "host-a"),
            {"ip:10.167.89.35"},
        )
        self.assertEqual(
            inventory.build_match_keys("", "10.167.89.35（私有）/172.16.8.20（公网）", "host-b"),
            {"ip:10.167.89.35", "ip:172.16.8.20"},
        )
        self.assertEqual(
            inventory.build_match_key("", "10.167.89.35（私有）", "host-a"),
            "ip:10.167.89.35",
        )

    def test_online_unprotected_excludes_host_when_any_ip_matches_exclusion_list(self) -> None:
        rows = [
            {
                "服务器名称": "host-multi-ip",
                "服务器ID": "",
                "IP地址": "10.167.89.35（私有）/172.16.8.20（公网）",
                "配额ID": "quota-1",
                "服务器状态": "运行中",
                "Agent状态": "在线",
                "风险状态": "高危",
                "防护状态": "未防护",
                "操作系统": "Linux",
                "版本类型": "正式",
                "企业项目": "项目A",
                "来源": "CMDB",
            }
        ]

        with TemporaryDirectory() as temp_dir:
            export_dir = Path(temp_dir)
            with (
                patch.object(inventory, "read_table_file", return_value=rows),
                patch.object(inventory, "get_owner_mapping_dict", return_value={"项目A": "张三"}),
                patch.object(
                    inventory,
                    "get_exclusion_match_keys",
                    side_effect=lambda key: (
                        {"ip:172.16.8.20"}
                        if key == "unquota-hosts"
                        else ({"id:container-snapshot"} if key == "unprotected-container-nodes" else set())
                    ),
                ),
                patch.object(inventory, "EXPORT_DIR", export_dir),
                patch.object(inventory, "create_result_history"),
                patch.object(inventory, "save_import_history"),
            ):
                result = inventory.generate_from_asset_file(export_dir / "asset.xlsx", "系统管理员")

        self.assertEqual(result["counts"]["online_unprotected"], 0)

    def test_protection_interrupted_requires_running_server(self) -> None:
        rows = [
            {
                "服务器名称": "host-interrupted",
                "服务器ID": "srv-1",
                "IP地址": "10.0.0.1",
                "配额ID": "quota-1",
                "服务器状态": "运行中",
                "Agent状态": "在线",
                "风险状态": "高危",
                "防护状态": "防护中断",
                "操作系统": "Linux",
                "版本类型": "正式",
                "企业项目": "项目A",
                "来源": "CMDB",
            },
            {
                "服务器名称": "host-no-agent",
                "服务器ID": "srv-2",
                "IP地址": "10.0.0.2",
                "配额ID": "quota-2",
                "服务器状态": "运行中",
                "Agent状态": "未安装",
                "风险状态": "高危",
                "防护状态": "防护中断",
                "操作系统": "Linux",
                "版本类型": "正式",
                "企业项目": "项目A",
                "来源": "CMDB",
            },
            {
                "服务器名称": "host-not-running",
                "服务器ID": "srv-3",
                "IP地址": "10.0.0.3",
                "配额ID": "quota-3",
                "服务器状态": "已停止",
                "Agent状态": "在线",
                "风险状态": "高危",
                "防护状态": "防护中断",
                "操作系统": "Linux",
                "版本类型": "正式",
                "企业项目": "项目A",
                "来源": "CMDB",
            },
            {
                "服务器名称": "host-unprotected",
                "服务器ID": "srv-4",
                "IP地址": "10.0.0.4",
                "配额ID": "quota-4",
                "服务器状态": "运行中",
                "Agent状态": "在线",
                "风险状态": "高危",
                "防护状态": "未防护",
                "操作系统": "Linux",
                "版本类型": "正式",
                "企业项目": "项目A",
                "来源": "CMDB",
            },
        ]

        with TemporaryDirectory() as temp_dir:
            export_dir = Path(temp_dir)
            with (
                patch.object(inventory, "read_table_file", return_value=rows),
                patch.object(inventory, "get_owner_mapping_dict", return_value={"项目A": "张三"}),
                patch.object(
                    inventory,
                    "get_exclusion_match_keys",
                    side_effect=lambda key: {"id:container-snapshot"} if key == "unprotected-container-nodes" else set(),
                ),
                patch.object(inventory, "EXPORT_DIR", export_dir),
                patch.object(inventory, "create_result_history"),
                patch.object(inventory, "save_import_history"),
            ):
                result = inventory.generate_from_asset_file(export_dir / "asset.xlsx", "系统管理员")

        self.assertEqual(result["counts"]["protection_interrupted"], 2)
        self.assertEqual(
            [row["服务器ID"] for row in result["previews"]["protection_interrupted"]],
            ["srv-1", "srv-2"],
        )

    def test_generated_report_uses_byd_prefix_with_date_and_contains_summary_sheet(self) -> None:
        rows = [
            {
                "服务器名称": "host-a",
                "服务器ID": "srv-1",
                "IP地址": "10.0.0.1",
                "配额ID": "quota-1",
                "服务器状态": "运行中",
                "Agent状态": "在线",
                "风险状态": "高危",
                "防护状态": "未防护",
                "操作系统": "Linux",
                "版本类型": "正式",
                "企业项目": "项目A",
                "来源": "CMDB",
            },
            {
                "服务器名称": "host-b",
                "服务器ID": "srv-2",
                "IP地址": "10.0.0.2",
                "配额ID": "quota-2",
                "服务器状态": "运行中",
                "Agent状态": "在线",
                "风险状态": "高危",
                "防护状态": "未防护",
                "操作系统": "Linux",
                "版本类型": "正式",
                "企业项目": "项目A",
                "来源": "CMDB",
            },
        ]

        with TemporaryDirectory() as temp_dir:
            export_dir = Path(temp_dir)
            with (
                patch.object(inventory, "read_table_file", return_value=rows),
                patch.object(inventory, "get_owner_mapping_dict", return_value={"项目A": "张三"}),
                patch.object(
                    inventory,
                    "get_exclusion_match_keys",
                    side_effect=lambda key: {"id:container-snapshot"} if key == "unprotected-container-nodes" else set(),
                ),
                patch.object(inventory, "EXPORT_DIR", export_dir),
                patch.object(inventory, "create_result_history"),
                patch.object(inventory, "save_import_history"),
            ):
                inventory.generate_from_asset_file(export_dir / "asset.xlsx", "系统管理员")

            xlsx_path = next(export_dir.glob("*/比亚迪Agent在线未添加防护配置主机列表-*.xlsx"))
            self.assertRegex(xlsx_path.name, r"^比亚迪Agent在线未添加防护配置主机列表-\d{4}-\d{2}-\d{2}\.xlsx$")

            workbook = load_workbook(xlsx_path)
            self.assertEqual(workbook.sheetnames, ["汇总", "比亚迪Agent在线未添加防护配置主机列表"])
            summary_sheet = workbook["汇总"]
            table_sheet = workbook["比亚迪Agent在线未添加防护配置主机列表"]

            self.assertEqual(summary_sheet["A1"].value, "负责人")
            self.assertEqual(summary_sheet["B1"].value, "服务器ID计数")
            self.assertEqual(summary_sheet["A2"].value, "张三")
            self.assertEqual(summary_sheet["B2"].value, 2)
            self.assertEqual(summary_sheet["A3"].value, "合计")
            self.assertEqual(summary_sheet["B3"].value, 2)
            self.assertEqual(table_sheet["A1"].value, inventory.OUTPUT_COLUMNS[0])
            container_column = inventory.OUTPUT_COLUMNS.index("是否为容器节点") + 1
            self.assertEqual(table_sheet.cell(row=1, column=container_column).value, "是否为容器节点")

    def test_generated_rows_mark_container_nodes_by_version_or_imported_snapshot(self) -> None:
        rows = [
            {
                "服务器名称": "host-container-version",
                "服务器ID": "srv-version",
                "IP地址": "10.0.0.1",
                "配额ID": "quota-1",
                "服务器状态": "运行中",
                "Agent状态": "在线",
                "风险状态": "高危",
                "防护状态": "未防护",
                "操作系统": "Linux",
                "版本类型": "容器版",
                "企业项目": "项目A",
                "来源": "CMDB",
            },
            {
                "服务器名称": "host-imported-container",
                "服务器ID": "srv-imported",
                "IP地址": "10.0.0.2",
                "配额ID": "quota-2",
                "服务器状态": "运行中",
                "Agent状态": "在线",
                "风险状态": "高危",
                "防护状态": "未防护",
                "操作系统": "Linux",
                "版本类型": "正式",
                "企业项目": "项目A",
                "来源": "CMDB",
            },
            {
                "服务器名称": "host-standard",
                "服务器ID": "srv-standard",
                "IP地址": "10.0.0.3",
                "配额ID": "quota-3",
                "服务器状态": "运行中",
                "Agent状态": "在线",
                "风险状态": "高危",
                "防护状态": "未防护",
                "操作系统": "Linux",
                "版本类型": "正式",
                "企业项目": "项目A",
                "来源": "CMDB",
            },
        ]

        with TemporaryDirectory() as temp_dir:
            export_dir = Path(temp_dir)
            with (
                patch.object(inventory, "read_table_file", return_value=rows),
                patch.object(inventory, "get_owner_mapping_dict", return_value={"项目A": "张三"}),
                patch.object(
                    inventory,
                    "get_exclusion_match_keys",
                    side_effect=lambda key: {"id:srv-imported"} if key == "unprotected-container-nodes" else set(),
                ),
                patch.object(inventory, "EXPORT_DIR", export_dir),
                patch.object(inventory, "create_result_history"),
                patch.object(inventory, "save_import_history"),
            ):
                result = inventory.generate_from_asset_file(export_dir / "asset.xlsx", "系统管理员")

        flags = {row["服务器ID"]: row["是否为容器节点"] for row in result["previews"]["online_unprotected"]}
        self.assertEqual(flags["srv-version"], "是")
        self.assertEqual(flags["srv-imported"], "是")
        self.assertEqual(flags["srv-standard"], "否")

    def test_generation_requires_imported_container_node_snapshot(self) -> None:
        rows = [
            {
                "服务器名称": "host-a",
                "服务器ID": "srv-1",
                "IP地址": "10.0.0.1",
                "配额ID": "quota-1",
                "服务器状态": "运行中",
                "Agent状态": "在线",
                "风险状态": "高危",
                "防护状态": "未防护",
                "操作系统": "Linux",
                "版本类型": "正式",
                "企业项目": "项目A",
                "来源": "CMDB",
            }
        ]

        with (
            patch.object(inventory, "read_table_file", return_value=rows),
            patch.object(inventory, "get_owner_mapping_dict", return_value={"项目A": "张三"}),
            patch.object(inventory, "get_exclusion_match_keys", return_value=set()),
        ):
            with self.assertRaisesRegex(ValueError, "请先导入未防护容器节点清单"):
                inventory.generate_from_asset_file(Path("asset.xlsx"), "系统管理员")
