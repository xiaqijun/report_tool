import json
from datetime import date, datetime, time as time_value, timedelta
from pathlib import Path
from urllib.parse import urlencode, urlparse
from uuid import uuid4

import requests
from openpyxl import load_workbook

from app import db


SETTINGS_KEY = "tencent_docs_settings"
TENCENT_DOCS_BASE_URL = "https://docs.qq.com"
MAX_RANGE_ROWS = 1000
MAX_RANGE_CELLS = 10000
MAX_CLEAR_COLUMNS = 200

HISTORY_DOCUMENTS = {
    "online_unprotected": (
        "online_unprotected_path",
        "tencent_online_unprotected_url",
        "在线未防护",
        "未添加防护配额信息",
    ),
    "agent_missing": (
        "agent_missing_path",
        "tencent_agent_missing_url",
        "Agent 未安装",
        "未安装Agent信息",
    ),
    "protection_interrupted": (
        "protection_interrupted_path",
        "tencent_protection_interrupted_url",
        "防护中断",
        "Agent防护中断信息",
    ),
}


def get_settings() -> dict[str, object]:
    raw_value = db.get_app_setting(SETTINGS_KEY)
    if not raw_value:
        return {}
    try:
        parsed = json.loads(raw_value)
    except json.JSONDecodeError:
        return {}
    return parsed if isinstance(parsed, dict) else {}


def get_public_settings() -> dict[str, object]:
    settings = get_settings()
    return {
        "client_id": settings.get("client_id", ""),
        "redirect_uri": settings.get("redirect_uri", ""),
        "target_document_url": settings.get("target_document_url", ""),
        "has_client_secret": bool(settings.get("client_secret")),
        "authorized": bool(settings.get("access_token") and settings.get("open_id")),
        "token_expires_at": settings.get("token_expires_at", ""),
    }


def save_configuration(payload: dict[str, object]) -> dict[str, object]:
    settings = get_settings()
    for key in ("client_id", "redirect_uri", "target_document_url", "token_expires_at"):
        if key in payload:
            value = str(payload.get(key, "") or "").strip()
            if key == "target_document_url" and value:
                _encoded_id, value = _parse_target_document_url(value)
            if key == "token_expires_at" and value:
                try:
                    expires_at = datetime.fromisoformat(value)
                except ValueError as error:
                    raise ValueError("令牌有效期格式应为 YYYY-MM-DD HH:MM:SS。") from error
                if expires_at.tzinfo is not None:
                    raise ValueError("令牌有效期请填写本地时间，不要包含时区。")
                value = expires_at.isoformat(timespec="seconds")
            settings[key] = value

    for key in ("client_secret", "access_token", "open_id"):
        value = str(payload.get(key, "") or "").strip()
        if value:
            settings[key] = value
            if key == "access_token":
                settings.pop("refresh_token", None)

    _save_settings(settings)
    return get_public_settings()


def build_authorize_url(state: str) -> str:
    settings = _require_configuration(require_secret=False)
    redirect_uri = str(settings["redirect_uri"])
    if not redirect_uri.lower().startswith("https://"):
        raise ValueError("腾讯文档 OAuth 回调地址必须使用 HTTPS。")
    query = urlencode(
        {
            "client_id": settings["client_id"],
            "redirect_uri": redirect_uri,
            "new_login": "1",
            "response_type": "code",
            "scope": "all",
            "state": state,
        }
    )
    return f"{TENCENT_DOCS_BASE_URL}/oauth/v2/authorize?{query}"


def exchange_authorization_code(code: str) -> dict[str, object]:
    settings = _require_configuration()
    token_data = _request_json(
        "GET",
        "/oauth/v2/token",
        params={
            "client_id": settings["client_id"],
            "client_secret": settings["client_secret"],
            "redirect_uri": settings["redirect_uri"],
            "grant_type": "authorization_code",
            "code": code,
        },
    )
    _store_token_response(settings, token_data, preserve_refresh_token=False)
    return get_public_settings()


def sync_history_documents(batch_code: str) -> dict[str, object]:
    history = db.get_result_history(batch_code)
    if history is None:
        raise ValueError("未找到对应的生成历史记录。")

    settings = _ensure_access_token()
    headers = _auth_headers(settings)
    target_document_url = str(settings.get("target_document_url", "") or "").strip()
    encoded_id, normalized_url = _parse_target_document_url(target_document_url)
    book_id = _convert_document_id(encoded_id, headers)
    sheets_by_title = _get_sheets_by_title(book_id, headers)
    links: dict[str, str] = {}
    pending_updates: list[tuple[str, dict[str, object], list[list[str]]]] = []

    for result_key, (path_field, _url_field, title, target_sheet_title) in HISTORY_DOCUMENTS.items():
        file_path = Path(str(history.get(path_field, "") or ""))
        if not file_path.is_file():
            raise ValueError(f"{title}文件不存在：{file_path}")
        sheet_info = sheets_by_title.get(target_sheet_title)
        if sheet_info is None:
            raise ValueError(f"目标腾讯表格缺少工作表：{target_sheet_title}")
        values = _read_detail_sheet_values(file_path)
        pending_updates.append((result_key, sheet_info, values))

    for result_key, sheet_info, values in pending_updates:
        updated_sheet_info = _replace_sheet_values(book_id, sheet_info, values, headers)
        sheet_id = str(updated_sheet_info["sheetID"])
        links[result_key] = f"{normalized_url}?tab={sheet_id}"

    db.update_result_history_tencent_docs(
        batch_code,
        links["online_unprotected"],
        links["agent_missing"],
        links["protection_interrupted"],
    )
    return {
        "batch_code": batch_code,
        "target_document_url": normalized_url,
        "links": links,
        "count": len(links),
    }


def _parse_target_document_url(value: str) -> tuple[str, str]:
    if not value:
        raise ValueError("请先配置目标腾讯表格链接。")
    parsed = urlparse(value)
    path_parts = [part for part in parsed.path.split("/") if part]
    if parsed.scheme != "https" or parsed.hostname != "docs.qq.com":
        raise ValueError("目标文档必须是 https://docs.qq.com 的在线表格链接。")
    if len(path_parts) < 2 or path_parts[0] != "sheet" or not path_parts[1]:
        raise ValueError("目标文档链接格式不正确，应为腾讯在线表格 /sheet/ 链接。")
    encoded_id = path_parts[1]
    return encoded_id, f"https://docs.qq.com/sheet/{encoded_id}"


def _convert_document_id(encoded_id: str, headers: dict[str, str]) -> str:
    converted = _request_json(
        "GET",
        "/openapi/drive/v2/util/converter",
        headers=headers,
        params={"type": 2, "value": encoded_id},
    )
    book_id = str(_unwrap_data(converted).get("fileID", "") or "")
    if not book_id:
        raise ValueError("腾讯文档 ID 转换响应缺少 fileID。")
    return book_id


def _get_sheets_by_title(book_id: str, headers: dict[str, str]) -> dict[str, dict[str, object]]:
    response = _request_json(
        "GET",
        f"/openapi/sheetbook/v2/{book_id}/sheets-info",
        headers=headers,
    )
    raw_sheets = _unwrap_data(response).get("sheetData", [])
    if not isinstance(raw_sheets, list):
        raise ValueError("腾讯表格子表查询响应格式不正确。")
    sheets: dict[str, dict[str, object]] = {}
    for item in raw_sheets:
        if isinstance(item, dict) and item.get("title") and item.get("sheetID"):
            sheets[str(item["title"])] = item
    return sheets


def _read_detail_sheet_values(file_path: Path) -> list[list[str]]:
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    try:
        candidates = [sheet for sheet in workbook.worksheets if sheet.title.strip() != "汇总"]
        detail_sheet = max(candidates or workbook.worksheets, key=lambda sheet: sheet.max_row * sheet.max_column)
        values = [
            [_cell_to_text(value) for value in row]
            for row in detail_sheet.iter_rows(values_only=True)
        ]
    finally:
        workbook.close()

    while values and not any(cell for cell in values[-1]):
        values.pop()
    if not values:
        raise ValueError(f"报表没有可同步的数据：{file_path.name}")
    column_count = max(len(row) for row in values)
    return [row + [""] * (column_count - len(row)) for row in values]


def _replace_sheet_values(
    book_id: str,
    sheet_info: dict[str, object],
    values: list[list[str]],
    headers: dict[str, str],
) -> dict[str, object]:
    sheet_id = str(sheet_info["sheetID"])
    existing_rows = int(sheet_info.get("rowCount", 0) or 0)
    existing_columns = int(sheet_info.get("columnCount", 0) or 0)
    column_count = max(len(row) for row in values)
    row_count = len(values)
    if (
        row_count > existing_rows or column_count > existing_columns
    ) and not _sheet_range_exists(book_id, sheet_id, row_count, column_count, headers):
        return _rebuild_sheet(book_id, sheet_info, values, headers)

    _write_sheet_values(book_id, sheet_id, values, headers)

    if existing_rows > row_count:
        _clear_sheet_range(
            book_id,
            sheet_id,
            row_count + 1,
            existing_rows,
            1,
            max(existing_columns, column_count),
            headers,
        )
    if existing_columns > column_count:
        _clear_sheet_range(
            book_id,
            sheet_id,
            1,
            row_count,
            column_count + 1,
            existing_columns,
            headers,
        )
    return sheet_info


def _write_sheet_values(
    book_id: str,
    sheet_id: str,
    values: list[list[str]],
    headers: dict[str, str],
) -> None:
    column_count = max(len(row) for row in values)
    rows_per_request = max(1, min(MAX_RANGE_ROWS, MAX_RANGE_CELLS // column_count))
    end_column = _column_name(column_count)
    for offset in range(0, len(values), rows_per_request):
        chunk = values[offset : offset + rows_per_request]
        start_row = offset + 1
        end_row = offset + len(chunk)
        cell_range = f"{sheet_id}!A{start_row}:{end_column}{end_row}"
        _request_json(
            "PUT",
            f"/openapi/sheetbook/v2/{book_id}/values/{cell_range}",
            headers=headers,
            json_body={"values": chunk},
        )



def _sheet_range_exists(
    book_id: str,
    sheet_id: str,
    row_count: int,
    column_count: int,
    headers: dict[str, str],
) -> bool:
    end_column = _column_name(column_count)
    try:
        _request_json(
            "GET",
            f"/openapi/spreadsheet/v3/files/{book_id}/{sheet_id}/A{row_count}:{end_column}{row_count}",
            headers=headers,
        )
        return True
    except ValueError as error:
        if "range" in str(error).lower() and "invalid" in str(error).lower():
            return False
        raise


def _rebuild_sheet(
    book_id: str,
    sheet_info: dict[str, object],
    values: list[list[str]],
    headers: dict[str, str],
) -> dict[str, object]:
    title = str(sheet_info.get("title", "") or "").strip()
    old_sheet_id = str(sheet_info["sheetID"])
    row_count = len(values)
    column_count = max(len(row) for row in values)
    if not title:
        raise ValueError("腾讯表格工作表缺少标题，无法安全扩容。")
    if row_count * column_count > MAX_RANGE_CELLS or column_count > MAX_CLEAR_COLUMNS:
        raise ValueError("目标工作表需要扩容，但报表超过腾讯表格单工作表创建限制。")

    suffix = f"_同步临时_{uuid4().hex[:6]}"
    temporary_title = f"{title[: max(1, 31 - len(suffix))]}{suffix}"
    temporary_sheet = _add_sheet(book_id, temporary_title, row_count, column_count, headers)
    temporary_sheet_id = str(temporary_sheet["sheetID"])
    try:
        _write_sheet_values(book_id, temporary_sheet_id, values, headers)
    except Exception:
        _delete_sheet(book_id, temporary_sheet_id, headers)
        raise

    _delete_sheet(book_id, old_sheet_id, headers)
    try:
        replacement_sheet = _add_sheet(book_id, title, row_count, column_count, headers)
        _write_sheet_values(book_id, str(replacement_sheet["sheetID"]), values, headers)
    except Exception as error:
        raise ValueError(
            f"目标工作表扩容未完成，数据已保留在临时工作表“{temporary_title}”：{error}"
        ) from error

    try:
        _delete_sheet(book_id, temporary_sheet_id, headers)
    except Exception as error:
        raise ValueError(
            f"目标工作表已更新，但临时工作表“{temporary_title}”清理失败：{error}"
        ) from error
    return replacement_sheet


def _add_sheet(
    book_id: str,
    title: str,
    row_count: int,
    column_count: int,
    headers: dict[str, str],
) -> dict[str, object]:
    response = _request_json(
        "POST",
        f"/openapi/spreadsheet/v3/files/{book_id}/batchUpdate",
        headers=headers,
        json_body={
            "requests": [
                {
                    "addSheetRequest": {
                        "title": title,
                        "rowCount": row_count,
                        "columnCount": column_count,
                    }
                }
            ]
        },
    )
    responses = _unwrap_data(response).get("responses", [])
    if not isinstance(responses, list) or not responses:
        raise ValueError("腾讯表格新增工作表响应缺少 responses。")
    first_response = responses[0] if isinstance(responses[0], dict) else {}
    add_response = first_response.get("addSheetResponse", {})
    properties = add_response.get("properties", {}) if isinstance(add_response, dict) else {}
    sheet_id = str(properties.get("sheetId", "") or "") if isinstance(properties, dict) else ""
    if not sheet_id:
        raise ValueError("腾讯表格新增工作表响应缺少 sheetId。")
    return {
        "sheetID": sheet_id,
        "title": str(properties.get("title", title) or title),
        "rowCount": row_count,
        "columnCount": column_count,
    }


def _delete_sheet(book_id: str, sheet_id: str, headers: dict[str, str]) -> None:
    _request_json(
        "POST",
        f"/openapi/spreadsheet/v3/files/{book_id}/batchUpdate",
        headers=headers,
        json_body={"requests": [{"deleteSheetRequest": {"sheetId": sheet_id}}]},
    )


def _clear_sheet_range(
    book_id: str,
    sheet_id: str,
    start_row: int,
    end_row: int,
    start_column_number: int,
    end_column_number: int,
    headers: dict[str, str],
) -> None:
    if end_row < start_row or end_column_number < start_column_number:
        return
    column_count = end_column_number - start_column_number + 1
    row_count = end_row - start_row + 1
    for column_offset in range(0, column_count, MAX_CLEAR_COLUMNS):
        clear_columns = min(MAX_CLEAR_COLUMNS, column_count - column_offset)
        rows_per_request = max(1, min(MAX_RANGE_ROWS, MAX_RANGE_CELLS // clear_columns))
        start_column = _column_name(start_column_number + column_offset)
        end_column = _column_name(start_column_number + column_offset + clear_columns - 1)
        for row_offset in range(0, row_count, rows_per_request):
            chunk_start_row = start_row + row_offset
            chunk_end_row = min(end_row, chunk_start_row + rows_per_request - 1)
            cell_range = f"{sheet_id}!{start_column}{chunk_start_row}:{end_column}{chunk_end_row}"
            _request_json(
                "POST",
                f"/openapi/sheetbook/v2/{book_id}/values/{cell_range}:clear",
                headers=headers,
            )


def _column_name(column_number: int) -> str:
    if column_number <= 0:
        raise ValueError("表格列数必须大于 0。")
    letters = ""
    while column_number:
        column_number, remainder = divmod(column_number - 1, 26)
        letters = chr(65 + remainder) + letters
    return letters


def _cell_to_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, (date, time_value)):
        return value.isoformat()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    return str(value)


def _auth_headers(settings: dict[str, object]) -> dict[str, str]:
    return {
        "Access-Token": str(settings["access_token"]),
        "Client-Id": str(settings["client_id"]),
        "Open-Id": str(settings["open_id"]),
    }


def _ensure_access_token() -> dict[str, object]:
    settings = get_settings()
    missing = [
        key
        for key in ("client_id", "access_token", "open_id")
        if not str(settings.get(key, "") or "").strip()
    ]
    if missing:
        raise ValueError(f"腾讯文档访问配置不完整：{', '.join(missing)}")

    expires_at = str(settings.get("token_expires_at", "") or "")
    if expires_at:
        try:
            needs_refresh = datetime.fromisoformat(expires_at) <= datetime.now() + timedelta(minutes=5)
        except ValueError:
            needs_refresh = True
        if needs_refresh:
            settings = _refresh_access_token(settings)
    return settings


def _refresh_access_token(settings: dict[str, object]) -> dict[str, object]:
    refresh_token = str(settings.get("refresh_token", "") or "")
    if not refresh_token:
        raise ValueError("腾讯文档授权已过期，请重新授权。")
    token_data = _request_json(
        "GET",
        "/oauth/v2/token",
        params={
            "client_id": settings["client_id"],
            "client_secret": settings["client_secret"],
            "grant_type": "refresh_token",
            "refresh_token": refresh_token,
        },
    )
    return _store_token_response(settings, token_data, preserve_refresh_token=True)


def _store_token_response(
    settings: dict[str, object], token_data: dict[str, object], *, preserve_refresh_token: bool
) -> dict[str, object]:
    access_token = str(token_data.get("access_token", "") or "")
    open_id = str(token_data.get("user_id", token_data.get("open_id", "")) or "")
    if not access_token or not open_id:
        raise ValueError("腾讯文档授权响应缺少 Access Token 或 Open ID。")
    settings["access_token"] = access_token
    settings["open_id"] = open_id
    refresh_token = str(token_data.get("refresh_token", "") or "")
    if refresh_token or not preserve_refresh_token:
        settings["refresh_token"] = refresh_token
    expires_in = int(token_data.get("expires_in", 0) or 0)
    settings["token_expires_at"] = (datetime.now() + timedelta(seconds=max(expires_in, 0))).isoformat(timespec="seconds")
    _save_settings(settings)
    return settings


def _require_configuration(*, require_secret: bool = True) -> dict[str, object]:
    settings = get_settings()
    required = ["client_id", "redirect_uri"]
    if require_secret:
        required.append("client_secret")
    missing = [key for key in required if not str(settings.get(key, "") or "").strip()]
    if missing:
        raise ValueError(f"腾讯文档配置不完整：{', '.join(missing)}")
    return settings


def _request_json(
    method: str,
    path: str,
    *,
    params: dict[str, object] | None = None,
    data: dict[str, object] | None = None,
    json_body: dict[str, object] | None = None,
    headers: dict[str, str] | None = None,
) -> dict[str, object]:
    response = requests.request(
        method,
        f"{TENCENT_DOCS_BASE_URL}{path}",
        params=params,
        data=data,
        json=json_body,
        headers={"Accept": "application/json", **(headers or {})},
        timeout=30,
    )
    try:
        payload = response.json()
    except ValueError:
        response.raise_for_status()
        raise ValueError("腾讯文档接口返回了无法识别的数据。")
    if not isinstance(payload, dict):
        raise ValueError("腾讯文档接口返回了无法识别的数据。")
    error_code = payload.get("ret", payload.get("code", 0))
    if not response.ok or error_code not in (0, None):
        raise ValueError(str(payload.get("msg", payload.get("message", f"腾讯文档接口错误：{error_code}"))))
    return payload


def _unwrap_data(payload: dict[str, object]) -> dict[str, object]:
    data = payload.get("data")
    return data if isinstance(data, dict) else payload


def _save_settings(settings: dict[str, object]) -> None:
    db.save_app_setting(SETTINGS_KEY, json.dumps(settings, ensure_ascii=False))
