from __future__ import annotations

import hashlib
import os
import re
import shutil
import zipfile
from dataclasses import dataclass
from datetime import date, datetime, time as datetime_time
from pathlib import Path, PurePosixPath
from tempfile import TemporaryDirectory
from typing import Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


PROVENANCE_HEADERS = ("来源压缩包", "来源文件", "来源行号")
SUPPORTED_WORKBOOK_SUFFIXES = {".xlsx", ".xlsm"}
MAX_INPUT_FILES = 100
MAX_WORKBOOKS = 5000
MAX_ARCHIVE_UNCOMPRESSED_BYTES = 4 * 1024 * 1024 * 1024
EXCEL_MAX_ROWS = 1_048_576


@dataclass(frozen=True)
class TableInput:
    path: Path
    display_name: str


@dataclass(frozen=True)
class SourceWorkbook:
    path: Path
    source_archive: str
    source_file: str


def _encode_value(value: object) -> tuple[bytes, bytes]:
    if value is None:
        return b"N", b""
    if isinstance(value, bool):
        return b"B", b"1" if value else b"0"
    if isinstance(value, datetime):
        return b"D", value.isoformat(timespec="microseconds").encode("utf-8")
    if isinstance(value, date):
        return b"d", value.isoformat().encode("utf-8")
    if isinstance(value, datetime_time):
        return b"t", value.isoformat(timespec="microseconds").encode("utf-8")
    if isinstance(value, int):
        return b"I", str(value).encode("ascii")
    if isinstance(value, float):
        return b"F", value.hex().encode("ascii")
    if isinstance(value, bytes):
        return b"Y", value
    return b"S", str(value).encode("utf-8", errors="surrogatepass")


def _row_digest(values: Sequence[object]) -> bytes:
    digest = hashlib.blake2b(digest_size=16)
    for value in values:
        marker, payload = _encode_value(value)
        digest.update(marker)
        digest.update(len(payload).to_bytes(4, "big"))
        digest.update(payload)
    return digest.digest()


def _normalize_headers(values: Sequence[object]) -> tuple[str, ...]:
    headers = ["" if value is None else str(value).strip() for value in values]
    while headers and not headers[-1]:
        headers.pop()
    if not headers or not any(headers):
        raise ValueError("Excel 表头为空。")
    if any(header in PROVENANCE_HEADERS for header in headers):
        raise ValueError("源表格已包含来源追踪列，请移除后重新处理。")
    return tuple(headers)


def _collect_workbooks(inputs: Sequence[TableInput], extract_dir: Path) -> list[SourceWorkbook]:
    if not inputs:
        raise ValueError("请至少上传一个 Excel 或 ZIP 文件。")
    if len(inputs) > MAX_INPUT_FILES:
        raise ValueError(f"一次最多上传 {MAX_INPUT_FILES} 个文件。")

    workbooks: list[SourceWorkbook] = []
    extract_dir.mkdir(parents=True, exist_ok=True)
    extracted_bytes = 0

    for input_index, item in enumerate(inputs, start=1):
        suffix = Path(item.display_name).suffix.lower()
        if suffix in SUPPORTED_WORKBOOK_SUFFIXES:
            workbooks.append(SourceWorkbook(item.path, "", item.display_name))
            continue
        if suffix != ".zip":
            raise ValueError(f"不支持的文件类型：{item.display_name}")

        try:
            archive = zipfile.ZipFile(item.path)
        except zipfile.BadZipFile as error:
            raise ValueError(f"ZIP 文件损坏：{item.display_name}") from error

        with archive:
            members = sorted(
                (
                    info
                    for info in archive.infolist()
                    if not info.is_dir()
                    and not PurePosixPath(info.filename).name.startswith("~$")
                    and PurePosixPath(info.filename).suffix.lower() in SUPPORTED_WORKBOOK_SUFFIXES
                ),
                key=lambda info: info.filename.casefold(),
            )
            for member_index, info in enumerate(members, start=1):
                if info.flag_bits & 0x1:
                    raise ValueError(f"不支持加密的 ZIP 文件：{item.display_name}")
                extracted_bytes += int(info.file_size)
                if extracted_bytes > MAX_ARCHIVE_UNCOMPRESSED_BYTES:
                    raise ValueError("ZIP 解压后的 Excel 文件总量超过 4GB 限制。")
                if len(workbooks) >= MAX_WORKBOOKS:
                    raise ValueError(f"一次最多处理 {MAX_WORKBOOKS} 个工作簿。")

                member_suffix = PurePosixPath(info.filename).suffix.lower()
                extracted_path = extract_dir / f"{input_index:04d}-{member_index:05d}{member_suffix}"
                with archive.open(info) as source, extracted_path.open("wb") as target:
                    shutil.copyfileobj(source, target, length=1024 * 1024)
                workbooks.append(
                    SourceWorkbook(
                        extracted_path,
                        item.display_name,
                        PurePosixPath(info.filename).as_posix(),
                    )
                )

    if not workbooks:
        raise ValueError("上传内容中没有找到 .xlsx 或 .xlsm 文件。")
    return workbooks


def _append_styled_header(sheet, headers: Sequence[str]) -> None:
    cells = []
    for value in headers:
        cell = WriteOnlyCell(sheet, value=value)
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(color="FFFFFF", bold=True, size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cells.append(cell)
    sheet.append(cells)


def _configure_sheet(sheet, headers: Sequence[str]) -> None:
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    sheet.row_dimensions[1].height = 30
    for index, header in enumerate(headers, start=1):
        width = max(12, min(38, len(str(header)) * 2 + 6))
        if header == "来源文件":
            width = 42
        elif header == "来源压缩包":
            width = 30
        sheet.column_dimensions[get_column_letter(index)].width = width


def _verify_output(path: Path, expected_headers: Sequence[str], expected_rows: int) -> None:
    with zipfile.ZipFile(path) as archive:
        bad_entry = archive.testzip()
        if bad_entry is not None:
            raise RuntimeError(f"输出文件校验失败：{bad_entry}")

    row_pattern = re.compile(rb'<row r="(\d+)"')
    cell_pattern = re.compile(rb'<c r="([A-Z]+)(\d+)"')
    max_row = 0
    final_column = ""
    overlap = b""
    with zipfile.ZipFile(path) as archive:
        with archive.open("xl/worksheets/sheet1.xml") as sheet_file:
            while chunk := sheet_file.read(1024 * 1024):
                data = overlap + chunk
                for match in row_pattern.finditer(data):
                    max_row = max(max_row, int(match.group(1)))
                for match in cell_pattern.finditer(data):
                    row_number = int(match.group(2))
                    if row_number >= max_row:
                        final_column = match.group(1).decode("ascii")
                overlap = data[-128:]

    expected_final_column = get_column_letter(len(expected_headers))
    if max_row != expected_rows + 1 or final_column != expected_final_column:
        raise RuntimeError("输出工作表边界校验失败。")

    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        if workbook.sheetnames != ["合并明细"]:
            raise RuntimeError("输出工作表结构异常。")
        sheet = workbook["合并明细"]
        actual_headers = tuple(next(sheet.iter_rows(min_row=1, max_row=1, values_only=True)))
        if actual_headers != tuple(expected_headers):
            raise RuntimeError("输出表头校验失败。")
    finally:
        workbook.close()


def merge_table_files(
    inputs: Sequence[TableInput],
    output_path: Path,
    *,
    deduplicate: bool = True,
) -> dict[str, object]:
    """Stream compatible workbooks into one traced, optionally deduplicated workbook."""
    output_path = output_path.resolve()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    partial_path = output_path.with_name(f"{output_path.stem}.partial{output_path.suffix}")
    if partial_path.exists():
        partial_path.unlink()

    with TemporaryDirectory(prefix="table-tool-", dir=output_path.parent) as temp_dir:
        workbooks = _collect_workbooks(inputs, Path(temp_dir) / "extracted")
        output_book = Workbook(write_only=True)
        output_sheet = output_book.create_sheet("合并明细")
        canonical_headers: tuple[str, ...] | None = None
        seen: set[bytes] = set()
        rows_read = 0
        rows_written = 0
        duplicate_rows = 0

        try:
            for source in workbooks:
                try:
                    source_book = load_workbook(
                        source.path,
                        read_only=True,
                        data_only=True,
                        keep_links=False,
                    )
                except Exception as error:
                    raise ValueError(f"无法读取 Excel：{source.source_file}") from error

                try:
                    if len(source_book.sheetnames) != 1:
                        raise ValueError(f"每个 Excel 必须只有一个工作表：{source.source_file}")
                    source_sheet = source_book[source_book.sheetnames[0]]
                    rows = source_sheet.iter_rows(values_only=True)
                    try:
                        headers = _normalize_headers(tuple(next(rows)))
                    except StopIteration as error:
                        raise ValueError(f"Excel 为空：{source.source_file}") from error

                    if canonical_headers is None:
                        canonical_headers = headers
                        all_headers = (*canonical_headers, *PROVENANCE_HEADERS)
                        _configure_sheet(output_sheet, all_headers)
                        _append_styled_header(output_sheet, all_headers)
                    elif headers != canonical_headers:
                        raise ValueError(f"表头不一致：{source.source_file}")

                    column_count = len(canonical_headers)
                    for source_row_number, row in enumerate(rows, start=2):
                        values = tuple(row[:column_count])
                        if len(values) < column_count:
                            values += (None,) * (column_count - len(values))
                        if not any(value not in (None, "") for value in values):
                            continue
                        rows_read += 1
                        if deduplicate:
                            digest = _row_digest(values)
                            if digest in seen:
                                duplicate_rows += 1
                                continue
                            seen.add(digest)
                        if rows_written + 2 > EXCEL_MAX_ROWS:
                            raise ValueError("合并结果超过 Excel 单工作表 1,048,576 行限制。")
                        output_sheet.append(
                            (*values, source.source_archive, source.source_file, source_row_number)
                        )
                        rows_written += 1
                finally:
                    source_book.close()

            if canonical_headers is None:
                raise ValueError("没有找到可合并的表头。")
            final_column = get_column_letter(len(canonical_headers) + len(PROVENANCE_HEADERS))
            output_sheet.auto_filter.ref = f"A1:{final_column}{rows_written + 1}"
            output_sheet.sheet_properties.pageSetUpPr.fitToPage = True
            output_sheet.page_setup.fitToWidth = 1
            output_sheet.page_setup.fitToHeight = 0
            output_book.save(partial_path)
        except Exception:
            if not partial_path.exists():
                try:
                    output_book.save(partial_path)
                except Exception:
                    pass
            if partial_path.exists():
                partial_path.unlink()
            raise

    expected_headers = (*canonical_headers, *PROVENANCE_HEADERS)
    _verify_output(partial_path, expected_headers, rows_written)
    os.replace(partial_path, output_path)
    return {
        "input_files": len(inputs),
        "workbooks": len(workbooks),
        "rows_read": rows_read,
        "rows_written": rows_written,
        "duplicate_rows": duplicate_rows,
        "deduplicated": deduplicate,
        "columns": len(canonical_headers),
    }
